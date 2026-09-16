import io
import json
import unittest
import zipfile
from unittest.mock import patch

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from backend.document_parser import DocumentParseError, parse_document, validate_filename


def pdf_bytes(page_texts, *, encrypted=False):
    writer = PdfWriter()
    for text in page_texts:
        page = writer.add_blank_page(width=300, height=300)
        if text is not None:
            font = DictionaryObject({
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            })
            page[NameObject("/Resources")] = DictionaryObject({
                NameObject("/Font"): DictionaryObject({NameObject("/F1"): font}),
            })
            stream = DecodedStreamObject()
            stream.set_data(f"BT /F1 12 Tf 20 240 Td ({text}) Tj ET".encode("ascii"))
            page[NameObject("/Contents")] = stream
    if encrypted:
        writer.encrypt("sample-password")
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def docx_bytes():
    document = Document()
    document.add_heading("Policy", level=1)
    document.add_paragraph("First paragraph.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Item"
    table.cell(0, 1).text = "Decision"
    table.cell(1, 0).text = "Refund"
    table.cell(1, 1).text = "Allowed"
    document.add_paragraph("After the table.")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


class DocumentParserTests(unittest.TestCase):
    def assert_error(self, code, filename, raw):
        with self.assertRaises(DocumentParseError) as context:
            parse_document(filename, raw)
        self.assertEqual(context.exception.status_code, code)
        return str(context.exception)

    def test_plain_text_preserves_unicode_and_normalizes_line_endings(self):
        raw = "\ufeff결정 😀\r\n두 번째 줄\r마지막 줄\n".encode("utf-8")
        result = parse_document("C:\\fakepath\\회의.TXT", raw)
        self.assertEqual(result["body"], "결정 😀\n두 번째 줄\n마지막 줄")
        self.assertEqual(result["title"], "회의.TXT")
        self.assertIsNone(result["timestamp"])
        self.assertEqual(result["metadata"]["original_filename"], "회의.TXT")
        self.assertEqual(result["metadata"]["mime_type"], "text/plain")
        self.assertEqual(result["metadata"]["size_bytes"], len(raw))
        self.assertEqual(result["metadata"]["character_count"], len(result["body"]))
        self.assertEqual(result["metadata"]["line_count"], 3)
        self.assertEqual(result["metadata"]["offset_unit"], "unicode_code_points")

    def test_markdown_keeps_headings_links_and_lists(self):
        content = "# Policy\n\n- [Reference](https://example.com/policy)\n- Confirmed"
        result = parse_document("policy.md", content.encode())
        self.assertEqual(result["body"], content)
        self.assertEqual(result["metadata"]["mime_type"], "text/markdown")
        self.assertEqual(result["metadata"]["blocks"][0]["kind"], "markdown")

    def test_filename_removes_browser_paths_and_rejects_unsafe_names(self):
        self.assertEqual(validate_filename("../../notes.md"), "notes.md")
        self.assertEqual(validate_filename("C:\\fakepath\\notes.md"), "notes.md")
        for name in ("", None, "\x00notes.md", "bad\nname.txt", "../", "a:notes.md", "a?b.txt", "a" * 256 + ".txt", "\ud800.txt"):
            with self.subTest(name=name):
                self.assert_error(400, name, b"Text")
        self.assert_error(415, "document.html", b"Text")
        self.assert_error(415, "document", b"Text")

    def test_empty_and_non_utf8_text_are_rejected(self):
        self.assert_error(400, "empty.txt", b"")
        self.assert_error(422, "blank.txt", b" \n\t")
        self.assert_error(422, "binary.txt", b"\xff\xfea\x00")
        self.assert_error(422, "null.txt", b"some\x00text")

    def test_raw_and_extracted_text_limits_reject_instead_of_truncating(self):
        with patch("backend.document_parser.MAX_FILE_BYTES", 3):
            self.assert_error(413, "large.txt", b"1234")
        with patch("backend.document_parser.MAX_BODY_CHARS", 5):
            self.assert_error(413, "large.txt", b"123456")
            self.assertEqual(parse_document("exact.txt", b"12345")["body"], "12345")

    def test_pdf_preserves_page_numbers_and_offsets(self):
        result = parse_document("policy.pdf", pdf_bytes(["First decision", "Second decision"]))
        self.assertEqual(result["metadata"]["page_count"], 2)
        self.assertFalse(result["metadata"]["ocr_performed"])
        for page, expected in zip(result["metadata"]["pages"], ("First decision", "Second decision")):
            self.assertEqual(result["body"][page["start"]:page["end"]], expected)
            self.assertTrue(page["has_text"])
        self.assertEqual([page["page"] for page in result["metadata"]["pages"]], [1, 2])
        self.assertNotIn("warnings", result["metadata"])

    def test_pdf_blank_page_records_warning_without_inventing_text(self):
        result = parse_document("mixed.pdf", pdf_bytes(["Decision", None, "Evidence"]))
        blank = result["metadata"]["pages"][1]
        self.assertEqual(blank["page"], 2)
        self.assertFalse(blank["has_text"])
        self.assertEqual(blank["start"], blank["end"])
        self.assertEqual(result["metadata"]["warnings"][0]["pages"], [2])
        self.assertEqual(result["body"], "Decision\n\nEvidence")

    def test_pdf_without_text_is_rejected_with_ocr_explanation(self):
        message = self.assert_error(422, "scan.pdf", pdf_bytes([None]))
        self.assertIn("OCR", message)

    def test_pdf_encrypted_and_malformed_documents_are_rejected(self):
        self.assert_error(422, "broken.pdf", b"Not a PDF")
        self.assert_error(422, "truncated.pdf", b"%PDF-1.7\ntruncated")
        self.assert_error(422, "private.pdf", pdf_bytes(["Secret"], encrypted=True))

    def test_pdf_page_and_stream_limits_are_enforced(self):
        with patch("backend.document_parser.MAX_PDF_PAGES", 1):
            self.assert_error(413, "pages.pdf", pdf_bytes(["One", "Two"]))
        with patch("backend.document_parser.MAX_PDF_STREAM_BYTES", 3):
            self.assert_error(413, "stream.pdf", pdf_bytes(["Text"]))
        with patch("backend.document_parser.MAX_PDF_TOTAL_STREAM_BYTES", 3):
            self.assert_error(413, "total.pdf", pdf_bytes(["Text"]))

    def test_docx_keeps_headings_paragraphs_and_table_in_document_order(self):
        result = parse_document("policy.docx", docx_bytes())
        self.assertEqual(result["body"], "Policy\n\nFirst paragraph.\n\nItem\tDecision\nRefund\tAllowed\n\nAfter the table.")
        blocks = result["metadata"]["blocks"]
        self.assertEqual([block["kind"] for block in blocks], ["heading", "paragraph", "table", "paragraph"])
        self.assertEqual(blocks[0]["level"], 1)
        self.assertEqual((blocks[2]["rows"], blocks[2]["columns"]), (2, 2))
        self.assertEqual(result["body"][blocks[2]["start"]:blocks[2]["end"]], "Item\tDecision\nRefund\tAllowed")
        self.assertIsNone(result["timestamp"])

    def test_docx_nested_table_and_merged_cells_keep_text(self):
        document = Document()
        table = document.add_table(rows=1, cols=2)
        cell = table.cell(0, 0).merge(table.cell(0, 1))
        cell.text = "Before"
        nested = cell.add_table(rows=1, cols=1)
        nested.cell(0, 0).text = "Nested decision"
        cell.add_paragraph("After")
        output = io.BytesIO()
        document.save(output)
        body = parse_document("table.docx", output.getvalue())["body"]
        self.assertEqual(body.count("Nested decision"), 1)
        self.assertLess(body.index("Before"), body.index("Nested decision"))
        self.assertLess(body.index("Nested decision"), body.index("After"))

    def test_docx_invalid_empty_and_oversized_package_are_rejected(self):
        self.assert_error(422, "broken.docx", b"Not a ZIP")
        document = Document()
        output = io.BytesIO()
        document.save(output)
        self.assert_error(422, "empty.docx", output.getvalue())
        with patch("backend.document_parser.MAX_DOCX_UNPACKED_BYTES", 100):
            self.assert_error(413, "large.docx", docx_bytes())
        with patch("backend.document_parser.MAX_DOCX_ENTRIES", 1):
            self.assert_error(413, "entries.docx", docx_bytes())
        output = io.BytesIO()
        with zipfile.ZipFile(output, "w") as package:
            package.writestr("unrelated.txt", "Hello")
        self.assert_error(422, "unrelated.docx", output.getvalue())

    def test_json_scalar_paths_are_readable_and_resolve_to_body_offsets(self):
        source = {"project": "Alpha", "decisions": [{"status": "approved", "count": 2, "active": True}], "owner.name": "Lee", "unknown": None}
        result = parse_document("project.json", json.dumps(source).encode())
        self.assertIn("$.project: Alpha", result["body"])
        self.assertIn("$.decisions[0].status: approved", result["body"])
        self.assertIn("$.decisions[0].count: 2", result["body"])
        self.assertIn("$.decisions[0].active: true", result["body"])
        self.assertIn('$["owner.name"]: Lee', result["body"])
        self.assertIn("$.unknown: null", result["body"])
        for block in result["metadata"]["blocks"]:
            self.assertTrue(result["body"][block["start"]:block["end"]].startswith(block["path"] + ": "))
        self.assertIsNone(result["timestamp"])

    def test_json_root_scalar_and_utf8_bom_are_supported(self):
        result = parse_document("scalar.json", b'\xef\xbb\xbf"A decision"')
        self.assertEqual(result["body"], "$: A decision")
        self.assertEqual(result["metadata"]["blocks"][0]["path"], "$")

    def test_json_invalid_duplicate_nonfinite_and_empty_values_are_rejected(self):
        for raw in (b"{", b'{"key": 1, "key": 2}', b'{"nested": {"key": 1, "key": 2}}', b"NaN", b"Infinity", b"-Infinity", b"1e9999", b"{}", b"[]", b'""', b'"\\ud800"'):
            with self.subTest(raw=raw):
                self.assert_error(422, "invalid.json", raw)

    def test_json_depth_node_and_output_limits_are_enforced(self):
        with patch("backend.document_parser.MAX_JSON_DEPTH", 2):
            self.assert_error(413, "deep.json", b'{"a":{"b":{"c":1}}}')
        with patch("backend.document_parser.MAX_JSON_NODES", 3):
            self.assert_error(413, "items.json", b"[1,2,3]")
        with patch("backend.document_parser.MAX_BODY_CHARS", 20):
            self.assert_error(413, "output.json", b'{"first":"value", "second":"value"}')


if __name__ == "__main__":
    unittest.main()
